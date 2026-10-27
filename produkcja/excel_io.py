"""Wczytywanie i zapis listy produkcyjnej (.xlsx).

Import: wartości komórek (z cache Excela) trafiają do bazy, formuły oryginału są
zapamiętane osobno. Eksport: ładuje wczytany oryginał i nadpisuje tylko komórki
zmienione w aplikacji, więc formatowanie, formuły i ukryte arkusze zostają.
"""
import json
import re
from datetime import datetime, date, time, timedelta

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

EPOCH = datetime(1899, 12, 30)
ISO_DATE = re.compile(r'^\d{4}-\d{2}-\d{2}$')
ISO_DATETIME = re.compile(r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}$')
WORKDAY_RE = re.compile(r'^=WORKDAY\(A\d+,\s*C\d+', re.I)


def normalize_cell(v):
    """Wartość komórki -> typ do JSON (str / int / float / None)."""
    if v is None:
        return None
    if isinstance(v, datetime):
        if v.year <= 1900:   # liczba sformatowana jako data przez Excel, wracamy do liczby
            return round((v - EPOCH).total_seconds() / 86400, 4)
        return v.strftime('%Y-%m-%d') if v.time() == time(0, 0) else v.strftime('%Y-%m-%d %H:%M')
    if isinstance(v, date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, time):
        return v.strftime('%H:%M')
    if isinstance(v, bool):
        return 'TAK' if v else 'NIE'
    if isinstance(v, (int, float)):
        return int(v) if isinstance(v, float) and v.is_integer() else v
    s = str(v).strip()
    return s or None


def parse_input(text):
    """Tekst wpisany w aplikacji -> wartość komórki (liczby i daty jako typy)."""
    if text is None:
        return None
    s = str(text).strip()
    if not s:
        return None
    if ISO_DATE.match(s) or ISO_DATETIME.match(s):
        return s
    m = re.match(r'^(\d{1,2})[.](\d{1,2})[.](\d{4})$', s)
    if m:
        return f'{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}'
    if re.match(r'^-?\d+$', s):
        return int(s)
    if re.match(r'^-?\d+[.,]\d+$', s):
        return float(s.replace(',', '.'))
    return s


def to_excel_value(v):
    """Wartość z bazy -> wartość do zapisu w komórce Excela."""
    if isinstance(v, str):
        if ISO_DATE.match(v):
            return datetime.strptime(v, '%Y-%m-%d')
        if ISO_DATETIME.match(v):
            return datetime.strptime(v, '%Y-%m-%d %H:%M')
    return v


def _norm_header(s):
    return re.sub(r'\s+', ' ', str(s or '')).strip().upper()


HEADERLESS_NAMES = ['ZAM', 'ROZP.', 'CZAS REALIZACJI', 'TERMIN REALIZACJI', 'NAZWA TOW.',
                    'ILOŚĆ', 'Z', 'ODBIORCA', 'NUMER ZO']

ROLE_LABELS = {
    'start': 'Data (rozpoczęcie, kolumna DATA)', 'order_date': 'Data zamówienia', 'lead_days': 'Czas realizacji (dni)', 'due': 'Termin realizacji',
    'product': 'Nazwa towaru', 'qty': 'Ilość', 'client': 'Odbiorca', 'zo': 'Numer ZO', 'color': 'Kolor',
}


def detect_mapping(columns, headerless):
    """Pola logiczne arkusza -> indeks kolumny, wg nagłówków (a bez nagłówków wg stałego układu)."""
    if headerless:
        return {'start': 1, 'order_date': 0, 'lead_days': 2, 'due': 3, 'product': 4, 'qty': 5, 'client': 7, 'zo': 8}
    heads = [_norm_header(c) for c in columns]
    m = {}

    def first(pred):
        for i, h in enumerate(heads):
            if h and pred(h):
                return i
        return None

    spec = {
        'order_date': lambda h: h in ('ZAM', 'ZAM.', 'DATA', 'DATA PRODUKCJI'),
        'lead_days':  lambda h: h.startswith('CZAS REA'),
        'start':      lambda h: h in ('DATA', 'ROZP.', 'ROZP', 'ROZ.', 'ROZ'),
        'due':        lambda h: h.startswith('TERMIN REAL'),
        'product':    lambda h: h.startswith('NAZWA') and 'ODBIORC' not in h,
        'qty':        lambda h: h in ('ILOŚĆ', 'ILOSC', 'SZT'),
        'client':     lambda h: 'ODBIORC' in h,
        'zo':         lambda h: h in ('NUMER ZO', 'ZO'),
        'color':      lambda h: h == 'KOLOR',
    }
    for key, pred in spec.items():
        idx = first(pred)
        if idx is not None:
            m[key] = idx
    return m


def _find_header_row(rows):
    """Wśród pierwszych wierszy szuka nagłówków (komórki tekstowe z NAZWA i ODBIORCA/ILOŚĆ)."""
    for i, r in enumerate(rows[:5]):
        heads = [_norm_header(c) for c in r if isinstance(c, str)]
        if (any(h.startswith('NAZWA') for h in heads)
                and any('ODBIORC' in h or h in ('ILOŚĆ', 'SZT') for h in heads)):
            return i
    return None


def _trim(lst):
    while lst and lst[-1] is None:
        lst.pop()
    return lst


DEFAULT_DAYS_BACK, DEFAULT_DAYS_FWD = 14, 14   # domyślny zakres dat widoku listy


def default_range(today=None):
    """Domyślny zakres dat widoku: od `DEFAULT_DAYS_BACK` dni wstecz do `DEFAULT_DAYS_FWD` naprzód."""
    today = today or date.today()
    return today - timedelta(days=DEFAULT_DAYS_BACK), today + timedelta(days=DEFAULT_DAYS_FWD)


def row_date(vals, mapping):
    """Data wiersza do filtra zakresu dat: kolumna DATA (rozpoczęcie, to widać w widoku), a gdy jej
    brak data zamówienia, potem termin realizacji. None gdy wiersz nie ma daty."""
    for key in ('start', 'order_date', 'due'):
        idx = mapping.get(key)
        v = vals[idx] if idx is not None and idx < len(vals) else None
        if isinstance(v, str) and ISO_DATE.match(v[:10]):
            try:
                return datetime.strptime(v[:10], '%Y-%m-%d').date()
            except ValueError:
                continue
    return None


def build_sheet(name, position, hidden, raw_v, raw_f=None):
    """Arkusz z surowych wierszy: nagłówki, mapowanie pól, wiersze danych i kolumny z formułami.
    `raw_f` (formuły) jest opcjonalne, dla CSV go nie ma."""
    raw_f = raw_f or []
    hdr = _find_header_row(raw_v)
    ncols = max((len(r) for r in raw_v), default=0)
    if hdr is not None:
        columns = []
        for j in range(ncols):
            h = raw_v[hdr][j] if j < len(raw_v[hdr]) else None
            columns.append(str(h).strip() if h not in (None, '') else get_column_letter(j + 1))
        start = hdr + 1
    else:
        columns = [get_column_letter(j + 1) for j in range(ncols)]
        # arkusz bez nagłówków (np. PROFIT): pierwsze kolumny mają ten sam układ co w
        # pozostałych arkuszach, resztę użytkownik nazwie w aplikacji
        for j, nm in enumerate(HEADERLESS_NAMES[:ncols]):
            columns[j] = nm
        start = 0
    mapping = detect_mapping(columns, headerless=hdr is None)
    rows, fcount = [], {}
    for i in range(start, len(raw_v)):
        vals = [normalize_cell(c) for c in raw_v[i]]
        forms = {}
        if i < len(raw_f):
            for j, c in enumerate(raw_f[i]):
                if isinstance(c, str) and c.startswith('=') and len(c) > 1:
                    forms[j] = c
                    fcount[j] = fcount.get(j, 0) + 1
        if not any(v is not None for v in vals) and not forms:
            continue
        rows.append({'src_row': i + 1, 'data': _trim(vals), 'formulas': forms or None})
    n = len(rows) or 1
    formula_cols = sorted(j for j, c in fcount.items() if c / n >= 0.3)
    return {
        'name': name, 'position': position, 'hidden': hidden,
        'header_row': hdr + 1 if hdr is not None else None,
        'columns': columns, 'mapping': mapping,
        'formula_cols': formula_cols, 'rows': rows,
    }


def read_workbook(path):
    """Zwraca listę arkuszy: dict(name, hidden, header_row, columns, mapping,
    formula_cols, rows=[dict(src_row, data, formulas)])."""
    wb_v = load_workbook(path, read_only=True, data_only=True)
    wb_f = load_workbook(path, read_only=True, data_only=False)
    sheets = []
    for pos, ws_v in enumerate(wb_v.worksheets):
        ws_f = wb_f[ws_v.title]
        raw_v = [list(r) for r in ws_v.iter_rows(values_only=True)]
        raw_f = [list(r) for r in ws_f.iter_rows(values_only=True)]
        if not any(c is not None for r in raw_v for c in r):
            continue
        sheets.append(build_sheet(ws_v.title, pos, ws_v.sheet_state != 'visible', raw_v, raw_f))
    wb_v.close()
    wb_f.close()
    return sheets


# ── CSV (eksport pojedynczego arkusza z Excela, UTF-8) ────────────────────────

def _csv_cell(text):
    """Komórka CSV -> wartość: puste = None, daty ISO/dd.mm.rrrr, liczby; numery z zerem
    na początku (np. 007) zostają tekstem."""
    s = (text or '').strip()
    if not s:
        return None
    m = re.match(r'^(\d{4}-\d{2}-\d{2}) 00:00(:00)?$', s)
    if m:
        return m.group(1)
    if re.match(r'^0\d+$', s):
        return s
    return parse_input(s)


def read_csv(data, name):
    """Wczytuje CSV (UTF-8, ewentualnie Windows-1250) jako jeden arkusz `name`."""
    import csv
    import io
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = data.decode('cp1250')
    lines = [ln for ln in text.splitlines()[:30] if ln.strip()]
    if not lines:
        raise ValueError('Pusty plik CSV.')
    delim = max(';,\t', key=lambda d: sum(ln.count(d) for ln in lines))
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    raw_v = [[_csv_cell(c) for c in row] for row in reader]
    if not any(c is not None for r in raw_v for c in r):
        raise ValueError('Plik CSV nie zawiera danych.')
    return build_sheet(name, 0, False, raw_v)


def _csv_out(v):
    if v is None:
        return ''
    if isinstance(v, float):
        return f'{v:g}'.replace('.', ',')
    return str(v)


def sheet_to_csv(sheet, rows):
    """CSV (UTF-8 z BOM, separator ;) z nagłówkami arkusza i jego wierszami w kolejności."""
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=';', lineterminator='\r\n')
    cols = sheet.columns
    w.writerow(cols)
    for r in rows:
        data = json.loads(r.data or '[]')
        data = data + [None] * (len(cols) - len(data))
        w.writerow([_csv_out(v) for v in data[:len(cols)]])
    return ('\ufeff' + buf.getvalue()).encode('utf-8')


# Nagłówki kolumn, w które pracownik wpisuje swoje inicjały (domyślny wybór etapów;
# kolumny z datami, np. D.zaw, pomijamy). Administrator może to zmienić w aplikacji.
STAGE_KEYWORDS = ('GIĘCIE', 'GIECIE', 'SPAW', 'DRZW', 'MASK', 'CZYSZ', 'MYJ', 'MAL',
                  'SKŁAD', 'ZŁOŻ', 'KJ', 'WYPAL', 'LASER', 'SZYK', 'GRZ', 'KWADR', 'WSADY')


def detect_stage_columns(columns, headerless):
    """Indeksy kolumn wyglądających na etapy produkcji (po nagłówkach)."""
    if headerless:
        return []
    out = []
    for i, c in enumerate(columns):
        h = _norm_header(c)
        if h and any(h.startswith(k) for k in STAGE_KEYWORDS):
            out.append(i)
    return out


def row_fields(data, mapping):
    """Pola logiczne wiersza (zo, produkt, klient...) z listy komórek."""
    out = {}
    for key, idx in mapping.items():
        out[key] = data[idx] if idx < len(data) else None
    return out


def build_search(data):
    parts = [str(v) for v in data[:14] if v is not None]
    return ' '.join(parts).lower()


def add_workdays(start_iso, days):
    """Termin = data + dni robocze (sob/nd pomijane; święta z arkusza liczy Excel po otwarciu pliku)."""
    d = datetime.strptime(start_iso[:10], '%Y-%m-%d').date()
    n = int(days)
    step = 1 if n >= 0 else -1
    while n != 0:
        d += timedelta(days=step)
        if d.weekday() < 5:
            n -= step
    return d.strftime('%Y-%m-%d')


def export_workbook(src_path, dest_path, sheets, only_sheet_ids=None):
    """Zapisuje kopię z wprowadzonymi zmianami.

    `sheets` to lista obiektów ProdSheet. Zmienione komórki są nadpisywane,
    usunięte wiersze czyszczone (pozycje wierszy zostają, by nie psuć formuł),
    nowe wiersze dopisywane na końcu z przesuniętymi formułami."""
    wb = load_workbook(src_path)   # formuły i formatowanie zachowane
    keep = set()
    for sh in sheets:
        if sh.name not in wb.sheetnames:
            continue
        if only_sheet_ids is not None and sh.id not in only_sheet_ids:
            continue
        keep.add(sh.name)
        ws = wb[sh.name]
        ncols = len(sh.columns)
        fcols = sh.formula_cols
        for r in sh.deleted_rows:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).value = None
        template_formulas = {}
        next_new = ws.max_row + 1   # wczytano tylko okno dat, więc nowe wiersze za końcem arkusza w pliku
        for row in sorted(sh.rows.all(), key=lambda x: x.position):
            data = json.loads(row.data or '[]')
            changed = json.loads(row.changed) if row.changed else []
            forms = json.loads(row.formulas) if row.formulas else {}
            if row.src_row is not None:
                for j, f in forms.items():
                    template_formulas[int(j)] = (row.src_row, f)
                for j in changed:
                    v = data[j] if j < len(data) else None
                    ws.cell(row=row.src_row, column=j + 1).value = to_excel_value(v)
            else:
                new_r = next_new
                next_new += 1
                for j, v in enumerate(data):
                    if v is not None:
                        ws.cell(row=new_r, column=j + 1).value = to_excel_value(v)
                for j in fcols:
                    tmpl = template_formulas.get(j)
                    f = forms.get(str(j))
                    if not f and tmpl:
                        col = get_column_letter(j + 1)
                        f = Translator(tmpl[1], origin=f'{col}{tmpl[0]}').translate_formula(f'{col}{new_r}')
                    if f:
                        ws.cell(row=new_r, column=j + 1).value = f
    if only_sheet_ids is not None:
        for name in list(wb.sheetnames):
            if name not in keep:
                del wb[name]
    try:
        wb.calculation.fullCalcOnLoad = True   # Excel przeliczy WORKDAY po otwarciu
    except Exception:
        pass
    wb.save(dest_path)

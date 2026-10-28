"""Testy modułu KJ Prestiż, etap 6: raporty (rejestr niezgodności, FPY, Pareto, koszt zawrotów, wyniki wg pracownika,
czas kontroli), filtry, eksport do Excela i karta kontroli zlecenia (ekran i PDF)."""
import io
from collections import Counter
from datetime import date, datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from app import app as flask_app, db
from prestiz_kj import demo_data, demo_history, reports_data as rd, service
from prestiz_kj.models import (PkjInspection, PkjInspectionMaterial, PkjOrder, PkjReturn, PkjReturnOperation)
from test_app import login, _csrf

SLUGS = ('niezgodnosci', 'fpy', 'pareto', 'koszt-zawrotow', 'pracownicy', 'czas-kontroli')


@pytest.fixture(autouse=True)
def _clean(app, tmp_path):
    flask_app.config['PKJ_UPLOAD_FOLDER'] = str(tmp_path / 'pkj')
    with app.app_context():
        service.reset_all()
    yield
    with app.app_context():
        service.reset_all()


@pytest.fixture
def seeded(app):
    with app.app_context():
        demo_data.seed_configuration()
        demo_data.seed_orders()
        assert demo_history.seed_history()


@pytest.fixture
def admin(client):
    login(client, 'admin', 'Admin1234!')
    return client


def _report(slug, **args):
    return rd.build_report(slug, args)


def _first_round_materials():
    return (PkjInspectionMaterial.query.join(PkjInspection).filter(PkjInspection.round == 1)
            .filter(PkjInspection.status.in_(['zakonczona', 'przerwana'])).all())


def _table(report, i=0):
    return report['tables'][i]


# ── Dostęp i renderowanie ─────────────────────────────────────────────────────

class TestAccessAndRender:
    def test_all_pages_render(self, admin, seeded):
        assert 'Raporty KJ Prestiż' in admin.get('/kj-prestiz/raporty').get_data(as_text=True)
        for slug in SLUGS:
            resp = admin.get(f'/kj-prestiz/raporty/{slug}')
            assert resp.status_code == 200, slug
            html = resp.get_data(as_text=True)
            assert 'Eksport' in html or 'Excel' in html
            assert 'Filtry:' in html and 'bez filtrów' in html
        assert admin.get('/kj-prestiz/raporty/nie-ma').status_code == 404
        assert admin.get('/kj-prestiz/raporty/nie-ma/xlsx').status_code == 404

    def test_requires_module_permission(self, client, seeded, app):
        login(client, 'oper', 'Oper1234!')
        with app.app_context():
            oid = PkjOrder.query.filter_by(number='ZP-DEMO-101').one().id
        urls = ['/kj-prestiz/raporty'] + [f'/kj-prestiz/raporty/{s}' for s in SLUGS] + [
            f'/kj-prestiz/raporty/{s}/xlsx' for s in SLUGS] + [f'/kj-prestiz/raporty/karta/{oid}',
                                                              f'/kj-prestiz/raporty/karta/{oid}.pdf']
        assert all(client.get(u).status_code == 403 for u in urls)

    def test_empty_database_renders(self, admin, app):
        with app.app_context():
            demo_data.seed_configuration()
        for slug in SLUGS:
            html = admin.get(f'/kj-prestiz/raporty/{slug}').get_data(as_text=True)
            assert 'Brak danych' in html, slug
            assert admin.get(f'/kj-prestiz/raporty/{slug}/xlsx').status_code == 200

    def test_chart_only_where_there_is_data(self, admin, seeded):
        for slug in ('fpy', 'pareto', 'koszt-zawrotow', 'czas-kontroli'):
            html = admin.get(f'/kj-prestiz/raporty/{slug}').get_data(as_text=True)
            assert 'id="pkj-chart"' in html and 'chart.js@4.4.4' in html and 'pkj-chart-data' in html
        for slug in ('niezgodnosci', 'pracownicy'):
            assert 'id="pkj-chart"' not in admin.get(f'/kj-prestiz/raporty/{slug}').get_data(as_text=True)


# ── FPY ───────────────────────────────────────────────────────────────────────

class TestFpy:
    def test_matches_direct_count(self, app, seeded):
        with app.app_context():
            mats = _first_round_materials()
            ok = sum(1 for m in mats if m.verdict == 'OK')
            rep = _report('fpy')
            total = _table(rep)['total']
            assert total[1] == len(mats) and total[2] == ok and total[3] == len(mats) - ok
            assert total[4] == round(100.0 * ok / len(mats), 1)
            assert rep['kpis'][0][1] == f'{total[4]:.1f}%'.replace('.', ',')

    def test_every_axis_adds_up_to_the_total(self, app, seeded):
        with app.app_context():
            expected = len(_first_round_materials())
            for axis in ('operacja', 'dzial', 'tydzien', 'material'):
                rep = _report('fpy', po=axis)
                rows = _table(rep)['rows']
                assert sum(r[1] for r in rows) == expected, axis
                assert all(r[1] == r[2] + r[3] for r in rows)
                assert all(0 <= r[4] <= 100 for r in rows)
                assert rep['chart']['labels'] == [r[0] for r in rows]
            assert _report('fpy', po='cokolwiek')['axis'] == 'operacja'

    def test_by_operation_follows_route_order_and_values(self, app, seeded):
        with app.app_context():
            rows = _table(_report('fpy'))['rows']
            assert [r[0] for r in rows] == ['Gięcie', 'Spawanie', 'Malowanie', 'Składanie']
            by_op = {r[0]: r for r in rows}
            spawanie = [m for m in _first_round_materials() if m.inspection_id in {
                i.id for i in PkjInspection.query.all() if i.order_operation_id and _op_name(i) == 'Spawanie'}]
            assert by_op['Spawanie'][1] == len(spawanie)
            assert by_op['Spawanie'][4] < by_op['Gięcie'][4] or by_op['Spawanie'][4] < 100

    def test_weekly_axis_is_chronological(self, app, seeded):
        with app.app_context():
            labels = _report('fpy', po='tydzien')['chart']['labels']
            assert labels and all(l.startswith('tydz. od ') for l in labels)
            assert len(labels) >= 4 and len(set(labels)) == len(labels)
            today = datetime.now(timezone.utc).date()
            days = []
            for l in labels:                                                  # "tydz. od dd.mm", rok wynika z "dziś"
                day, month = int(l[-5:-3]), int(l[-2:])
                d = date(today.year, month, day)
                days.append(d.replace(year=d.year - 1) if d > today + timedelta(days=1) else d)
            assert days == sorted(days)


def _op_name(insp):
    from prestiz_kj.models import PkjOrderOperation
    return db.session.get(PkjOrderOperation, insp.order_operation_id).operation_name


# ── Filtry ────────────────────────────────────────────────────────────────────

class TestFilters:
    def test_client_product_and_department(self, app, seeded):
        with app.app_context():
            all_total = _table(_report('fpy'))['total'][1]
            by_client = _table(_report('fpy', klient='Klient Alfa'))['total'][1]
            by_product = _table(_report('fpy', produkt='Szafa PSH 600x800x300'))['total'][1]
            by_dept = _table(_report('fpy', dzial='Spawanie'))['total'][1]
            assert 0 < by_client < all_total and 0 < by_product < all_total and 0 < by_dept < all_total
            clients = {r[0] for r in _table(_report('fpy', klient='Klient Alfa', po='dzial'))['rows']}
            assert clients                                   # dział wykonawcy zostaje jako oś
            assert _table(_report('fpy', klient='Nie ma takiego'))['rows'] == []
            both = _table(_report('fpy', klient='Klient Alfa', dzial='Spawanie'))['total'][1]
            assert both <= min(by_client, by_dept)

    def test_period(self, app, seeded):
        with app.app_context():
            today = datetime.now(timezone.utc).date()
            full = _table(_report('fpy'))['total'][1]
            recent = _table(_report('fpy', od=(today - timedelta(days=7)).isoformat()))['total'][1]
            old = _table(_report('fpy', do=(today - timedelta(days=15)).isoformat()))['total'][1]
            assert 0 < recent < full and 0 < old < full
            both = _table(_report('fpy', od=(today - timedelta(days=15)).isoformat(),
                                  do=(today - timedelta(days=8)).isoformat()))['total'][1]
            assert 0 < both < full
            assert _table(_report('fpy', od='2999-01-01'))['rows'] == []
            assert _table(_report('fpy', od='zła data', do='też'))['total'][1] == full        # błędne daty ignorowane

    def test_filters_text_and_query(self, app, seeded):
        with app.app_context():
            f = rd.parse_filters({'od': '2026-10-01', 'do': '2026-10-31', 'klient': 'K', 'dzial': 'Spawanie'})
            assert rd.describe_filters(f) == 'okres: 01.10.2026 do 31.10.2026, klient: K, dział: Spawanie'
            assert rd.filters_to_query(f) == {'od': '2026-10-01', 'do': '2026-10-31', 'klient': 'K',
                                              'dzial': 'Spawanie'}
            assert rd.describe_filters(rd.parse_filters({})) == 'bez filtrów (wszystkie dane)'

    def test_filter_form_keeps_values_and_options(self, admin, seeded):
        html = admin.get('/kj-prestiz/raporty/fpy?klient=Klient+Alfa&dzial=Spawanie&od=2026-01-01').get_data(as_text=True)
        assert 'value="Klient Alfa" selected' in html and 'value="Spawanie" selected' in html
        assert 'value="2026-01-01"' in html and 'okres: 01.01.2026' in html
        assert 'Klient Beta' in html and 'Szafa PSH 600x600x250' in html


# ── Rejestr niezgodności ──────────────────────────────────────────────────────

class TestRegister:
    def test_one_row_per_nok_material(self, app, seeded):
        with app.app_context():
            nok = PkjInspectionMaterial.query.filter_by(verdict='NOK').count()
            rows = _table(_report('niezgodnosci'))['rows']
            assert len(rows) == nok and nok >= 20
            dates = [r[0] for r in rows]
            assert dates == sorted(dates, reverse=True)
            assert all(r[1] and r[5] and r[6] for r in rows)          # zlecenie, materiał, punkt zawsze są
            assert all(r[7] for r in rows if r[11])                    # przyczyna tam, gdzie jest zawrót
            assert sum(1 for r in rows if r[11] == 'Złom') == 2
            assert any('przerwana' in r[6] for r in rows)
            assert any(r[10].startswith('sprawdzająca') for r in rows)
            statuses = Counter(r[13] for r in rows)
            assert statuses['Zamknięte'] > 10 and statuses['Wymagany zawrót'] == 0
            kpis = dict(_report('niezgodnosci')['kpis'])
            assert kpis['Niezgodności'] == str(nok) and kpis['Złomy'] == '2'

    def test_row_content_for_known_failure(self, app, seeded):
        with app.app_context():
            rows = _table(_report('niezgodnosci'))['rows']
            spoina = [r for r in rows if r[6] == 'Spoina narożna kompletna']
            assert len(spoina) >= 7
            r = spoina[0]
            assert r[4] == 'Spawanie' and r[5] == 'Drzwi' and r[7] == 'Spoina niepełna'
            assert r[2].startswith('Klient') and r[3].startswith('Szafa PSH')
            assert r[8] and r[9] and r[11] in ('Poprawka', 'Złom')
            assert r[12].startswith('ZN-ZP-DEMO-')

    def test_waiting_for_return_is_visible(self, app, seeded):
        with app.app_context():
            from prestiz_kj import inspections as insp_
            from prestiz_kj.models import PkjOrderMaterial
            o = PkjOrder.query.filter_by(number='ZP-DEMO-004').one()       # czeka na kontrolę Gięcia
            oop = [x for x in __import__('prestiz_kj.orders', fromlist=['x']).operations_of(o.id)
                   if x.operation_name == 'Gięcie'][0]
            from models import User
            insp = insp_.start_inspection(oop, [m.id for m in insp_.pending_materials(oop)],
                                          User.query.filter_by(username='admin').one())
            insp_.interrupt_inspection(insp, 't')
            db.session.commit()
            rows = [r for r in _table(_report('niezgodnosci'))['rows'] if r[1] == 'ZP-DEMO-004']
            assert len(rows) == 3 and all(r[13] == 'Wymagany zawrót' and r[7] is None for r in rows)


# ── Pareto ────────────────────────────────────────────────────────────────────

class TestPareto:
    def test_cause_axis(self, app, seeded):
        with app.app_context():
            rep = _report('pareto', os='przyczyna')
            rows = _table(rep)['rows']
            assert rows[0][0] == 'Spoina niepełna'
            assert [r[1] for r in rows] == sorted((r[1] for r in rows), reverse=True)
            assert rows[-1][3] == 100.0
            assert abs(sum(r[2] for r in rows) - 100) < 0.6
            assert all(a[3] <= b[3] for a, b in zip(rows, rows[1:]))
            assert rep['chart']['type'] == 'pareto' and rep['chart']['cumulative'][-1] == 100.0
            assert rep['chart']['max'] == 100
            assert _table(rep)['total'][1] == sum(r[1] for r in rows)

    def test_other_axes(self, app, seeded):
        with app.app_context():
            nok = PkjInspectionMaterial.query.filter_by(verdict='NOK').count()
            for axis in ('material', 'operacja'):
                rows = _table(_report('pareto', os=axis))['rows']
                assert sum(r[1] for r in rows) == nok, axis
            point = _table(_report('pareto'))['rows']
            assert point[0][0] == 'Spoina narożna kompletna'
            assert any(r[0] == 'Kontrola przerwana' for r in point)
            assert _report('pareto', os='xyz')['axis'] == 'punkt'

    def test_filter_changes_pareto(self, app, seeded):
        with app.app_context():
            rows = _table(_report('pareto', os='przyczyna', dzial='Gięcie'))['rows']
            assert rows and all(r[0] != 'Spoina niepełna' for r in rows)


# ── Koszt zawrotów ────────────────────────────────────────────────────────────

class TestCost:
    def test_totals_follow_assumptions(self, app, seeded):
        with app.app_context():
            rets = PkjReturn.query.all()
            ops = sum(PkjReturnOperation.query.filter_by(return_id=r.id).count() for r in rets)
            rep = _report('koszt-zawrotow')
            total = _table(rep)['total']
            assert total[1] == len(rets) and total[3] == 2 and total[2] == len(rets) - 2
            assert total[4] == ops
            assert total[5] == round(ops * 1.5, 1)
            assert total[6] == round(ops * 1.5 * 120, 2)
            assert total[7] == 700.0 and total[8] == round(total[6] + total[7], 2)
            assert total[9] is not None and 0 < total[9] < 30

    def test_settings_drive_the_cost(self, app, seeded):
        with app.app_context():
            service.set_setting('repair_hour_cost', '200')
            service.set_setting('repair_op_hours', '2')
            service.set_setting('scrap_material_cost', '500')
            db.session.commit()
            total = _table(_report('koszt-zawrotow'))['total']
            assert total[5] == total[4] * 2.0 and total[6] == total[5] * 200 and total[7] == 1000.0

    def test_axes(self, app, seeded):
        with app.app_context():
            n = PkjReturn.query.count()
            for axis in ('operacja', 'przyczyna', 'material', 'tydzien'):
                rep = _report('koszt-zawrotow', po=axis)
                rows = _table(rep)['rows']
                assert sum(r[1] for r in rows) == n, axis
                assert all(abs(r[8] - (r[6] + r[7])) < 0.01 for r in rows)
                assert rep['chart']['values'] == [r[8] for r in rows]
            assert [r[0] for r in _table(_report('koszt-zawrotow'))['rows']][0] == 'Gięcie'
            assert _report('koszt-zawrotow', po='zły')['axis'] == 'operacja'

    def test_open_repairs_have_no_elapsed_time_but_count(self, app, seeded):
        with app.app_context():
            rep = _report('koszt-zawrotow', po='przyczyna')
            assert 'Założenia' in rep['note'] and '1,5 h' in rep['note'].replace('1.5', '1,5')


# ── Pracownicy i czas kontroli ────────────────────────────────────────────────

class TestEmployeesAndTime:
    def test_employees(self, app, seeded):
        with app.app_context():
            rep = _report('pracownicy')
            rows = _table(rep)['rows']
            mats = [m for m in _first_round_materials()]
            assert _table(rep)['total'][2] == len(mats)
            assert _table(rep)['total'][3] == sum(1 for m in mats if m.verdict == 'NOK')
            names = {r[0] for r in rows}
            assert {'Cezary Lis', 'Dorota Kos'} <= names
            assert [r[3] for r in rows] == sorted((r[3] for r in rows), reverse=True)
            assert 'nie do kar' in rep['subtitle']
            only = _table(_report('pracownicy', dzial='Spawanie'))['rows']
            assert only and {r[1] for r in only} == {'Spawanie'}

    def test_inspection_time(self, app, seeded):
        with app.app_context():
            rep = _report('czas-kontroli')
            by_op, by_ctrl = rep['tables']
            finished = PkjInspection.query.filter_by(status='zakonczona').count()
            assert sum(r[1] for r in by_op['rows']) == finished == sum(r[1] for r in by_ctrl['rows'])
            assert all(0 < r[2] <= r[3] < 60 for r in by_op['rows'])
            assert {r[0] for r in by_ctrl['rows']} == {'Anna Kontroler', 'Piotr Kontroler'}
            assert rep['chart']['labels'] == [r[0] for r in by_op['rows']]


# ── Excel ─────────────────────────────────────────────────────────────────────

class TestExcel:
    def _wb(self, client, slug, qs=''):
        resp = client.get(f'/kj-prestiz/raporty/{slug}/xlsx{qs}')
        assert resp.status_code == 200
        assert resp.mimetype.endswith('spreadsheetml.sheet')
        assert f'kj-prestiz-{slug}.xlsx' in resp.headers['Content-Disposition']
        return load_workbook(io.BytesIO(resp.data))

    def test_fpy_workbook(self, admin, seeded, app):
        wb = self._wb(admin, 'fpy')
        ws = wb.active
        assert ws.title.startswith('FPY wg operacja') and ws['A1'].value == 'Dobrze za pierwszym razem (FPY)'
        assert 'bez filtrów' in ws['A2'].value
        assert [c.value for c in ws[4]][:5] == ['Grupa', 'Skontrolowane materiały', 'Zgodne za pierwszym razem',
                                                 'Niezgodne', 'FPY']
        assert ws['A5'].value == 'Gięcie' and isinstance(ws['B5'].value, int)
        assert 0 < ws['E5'].value <= 1 and ws['E5'].number_format == '0.0%'
        assert ws.freeze_panes == 'A5' and ws.auto_filter.ref.startswith('A4:E')
        last = ws.max_row
        assert any(ws.cell(row=r, column=1).value == 'Razem' for r in range(5, last + 1))

    def test_register_workbook_has_dates_and_all_rows(self, admin, seeded, app):
        ws = self._wb(admin, 'niezgodnosci').active
        with app.app_context():
            n = PkjInspectionMaterial.query.filter_by(verdict='NOK').count()
        data_rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] is not None and r[1]]
        assert len(data_rows) == n
        assert isinstance(data_rows[0][0], datetime) and ws['A5'].number_format == 'DD.MM.YYYY HH:MM'

    def test_cost_workbook_uses_money_and_two_sheets_for_time(self, admin, seeded):
        ws = self._wb(admin, 'koszt-zawrotow').active
        assert ws['G5'].number_format == '#,##0.00 "PLN"' and isinstance(ws['G5'].value, (int, float))
        wb = self._wb(admin, 'czas-kontroli')
        assert len(wb.sheetnames) == 2 and wb.sheetnames[0].startswith('Wg operacji')

    def test_workbook_respects_filters_and_axis(self, admin, seeded):
        wb = self._wb(admin, 'pareto', '?os=przyczyna&klient=Klient+Alfa')
        ws = wb.active
        assert 'klient: Klient Alfa' in ws['A2'].value and ws['A4'].value == 'Przyczyna'
        assert ws.title.startswith('Pareto wg przyczyna')

    def test_all_workbooks_open(self, admin, seeded):
        for slug in SLUGS:
            assert self._wb(admin, slug).sheetnames


# ── Karta kontroli ────────────────────────────────────────────────────────────

class TestCard:
    def _ids(self, app):
        with app.app_context():
            zp = PkjOrder.query.filter_by(number='ZP-DEMO-101').one()
            zn = PkjOrder.query.filter_by(parent_id=zp.id).first()
            empty = PkjOrder.query.filter_by(number='ZP-DEMO-002').one()
            return zp.id, zn.id if zn else None, empty.id

    def test_card_page(self, admin, seeded, app):
        zp, zn, empty = self._ids(app)
        html = admin.get(f'/kj-prestiz/raporty/karta/{zp}').get_data(as_text=True)
        assert 'Karta kontroli: ZP-DEMO-101' in html and 'Kontrola' in html and 'Wspólne' in html
        assert 'PDF' in html and 'Zawroty i zlecenia naprawcze' in html
        assert 'data-label="Punkt"' in html
        none = admin.get(f'/kj-prestiz/raporty/karta/{empty}').get_data(as_text=True)
        assert 'nie ma jeszcze żadnych kontroli' in none
        assert admin.get('/kj-prestiz/raporty/karta/99999').status_code == 404

    def test_repair_order_redirects_to_parent_card(self, admin, seeded, app):
        zp, zn, _ = self._ids(app)
        assert zn
        resp = admin.get(f'/kj-prestiz/raporty/karta/{zn}')
        assert resp.status_code == 302 and resp.headers['Location'].endswith(f'/karta/{zp}')
        assert admin.get(f'/kj-prestiz/raporty/karta/{zn}.pdf').data.startswith(b'%PDF')

    def test_pdf(self, admin, seeded, app):
        zp, _, empty = self._ids(app)
        resp = admin.get(f'/kj-prestiz/raporty/karta/{zp}.pdf')
        assert resp.status_code == 200 and resp.mimetype == 'application/pdf'
        assert resp.data.startswith(b'%PDF') and len(resp.data) > 5000
        assert 'karta-kontroli-ZP-DEMO-101.pdf' in resp.headers['Content-Disposition']
        assert admin.get(f'/kj-prestiz/raporty/karta/{empty}.pdf').data.startswith(b'%PDF')

    def test_pdf_pages_text_with_polish_characters(self, admin, seeded, app):
        zp, _, _ = self._ids(app)
        data = admin.get(f'/kj-prestiz/raporty/karta/{zp}.pdf').data
        try:
            from pypdf import PdfReader
        except ImportError:                                     # biblioteka tylko do weryfikacji
            pytest.skip('pypdf niedostępne')
        text = ' '.join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
        assert 'Karta kontroli zlecenia' in text and 'ZP-DEMO-101' in text and 'Spawanie' in text

    def test_reports_index_lists_cards(self, admin, seeded):
        html = admin.get('/kj-prestiz/raporty').get_data(as_text=True)
        assert 'ZP-DEMO-101' in html and '.pdf' in html and 'Rejestr niezgodności' in html
        assert 'Koszt zawrotów' in html and 'Czas kontroli' in html

    def test_card_link_on_order_page(self, admin, seeded, app):
        zp, _, empty = self._ids(app)
        assert 'Karta kontroli (PDF)' in admin.get(f'/kj-prestiz/zlecenia/{zp}').get_data(as_text=True)
        assert 'Karta kontroli (PDF)' not in admin.get(f'/kj-prestiz/zlecenia/{empty}').get_data(as_text=True)


# ── Formatowanie i ustawienia ─────────────────────────────────────────────────

class TestHelpers:
    def test_format_cell(self):
        assert rd.format_cell(None, 'int') == '-' and rd.format_cell('', 'text') == '-'
        assert rd.format_cell(3, 'int') == '3' and rd.format_cell(78.94, 'pct') == '78,9%'
        assert rd.format_cell(1234.5, 'money') == '1 234,50 PLN' and rd.format_cell(2.26, 'float1') == '2,3'
        assert rd.format_cell(datetime(2026, 10, 28, 10, 5), 'datetime') == '28.10.2026 11:05'      # czas zimowy
        assert rd.format_cell(date(2026, 1, 2), 'date') == '02.01.2026'

    def test_week_start_is_monday_in_local_time(self):
        assert rd.week_start(datetime(2026, 10, 28, 12)) == date(2026, 10, 26)
        assert rd.week_start(datetime(2026, 10, 25, 23, 30)) == date(2026, 10, 26)      # niedziela 23:30 UTC to już poniedziałek lokalnie

    def test_setting_for_repair_time(self, admin, app):
        with app.app_context():
            demo_data.seed_configuration()
        admin.post('/kj-prestiz/konfiguracja/ustawienia/zapisz-wszystko',
                   data={'repair_op_hours': '2,5', '_csrf_token': _csrf(admin)})
        with app.app_context():
            assert service.get_setting('repair_op_hours') == '2.5'
        assert 'Czas jednej operacji naprawczej' in admin.get('/kj-prestiz/konfiguracja/ustawienia').get_data(as_text=True)

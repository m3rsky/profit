"""Eksport raportów KJ Prestiż do Excela (openpyxl): jeden arkusz na tabelę, liczby jako liczby."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .reports_data import to_local

HEADER_FILL = PatternFill('solid', fgColor='1A3A5C')
TOTAL_FILL = PatternFill('solid', fgColor='EAF2F8')
THIN = Side(style='thin', color='D0D7E2')
FORMATS = {'int': '0', 'float1': '0.0', 'money': '#,##0.00 "PLN"', 'pct': '0.0%', 'datetime': 'DD.MM.YYYY HH:MM',
           'date': 'DD.MM.YYYY'}


def _cell_value(value, kind):
    if value is None:
        return None
    if kind == 'pct':
        return value / 100.0
    if kind == 'datetime':
        return to_local(value).replace(tzinfo=None)
    return value


def _sheet_name(title, used):
    name = ''.join(c for c in title if c not in '[]:*?/\\')[:31] or 'Arkusz'
    base, n = name, 2
    while name in used:
        name = f'{base[:28]} {n}'
        n += 1
    used.add(name)
    return name


def report_workbook(report):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for table in report['tables']:
        ws = wb.create_sheet(_sheet_name(table['title'], used))
        ws['A1'] = report['title']
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = 'Filtry: ' + report['filters_text'] + '. Dane demonstracyjne, moduł KJ Prestiż (pogląd).'
        ws['A2'].font = Font(italic=True, color='64748B')
        head_row = 4
        columns = table['columns']
        for c, (label, _) in enumerate(columns, 1):
            cell = ws.cell(row=head_row, column=c, value=label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        widths = [len(label) for label, _ in columns]
        rows = list(table['rows'])
        for r, row in enumerate(rows, head_row + 1):
            for c, ((_, kind), value) in enumerate(zip(columns, row), 1):
                cell = ws.cell(row=r, column=c, value=_cell_value(value, kind))
                if kind in FORMATS:
                    cell.number_format = FORMATS[kind]
                cell.border = Border(bottom=THIN)
                widths[c - 1] = max(widths[c - 1], len(str(cell.value if cell.value is not None else '')))
        last = head_row + len(rows)
        if table.get('total'):
            last += 1
            for c, ((_, kind), value) in enumerate(zip(columns, table['total']), 1):
                cell = ws.cell(row=last, column=c, value=_cell_value(value, kind))
                cell.font = Font(bold=True)
                cell.fill = TOTAL_FILL
                if kind in FORMATS:
                    cell.number_format = FORMATS[kind]
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 10), 48)
        if rows:
            ws.auto_filter.ref = f'A{head_row}:{get_column_letter(len(columns))}{head_row + len(rows)}'
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
        if report.get('note'):
            ws.cell(row=last + 2, column=1, value=report['note']).font = Font(italic=True, color='64748B')
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
